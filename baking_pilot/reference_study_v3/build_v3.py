from pathlib import Path
import json,shutil
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.oxml.xmlchemy import OxmlElement
from PIL import Image
P=Path(__file__).parent;BASE=P.parent.parent;A=P/'assets';R=P/'results';r=Presentation(BASE/'dynamic_object_baking_v2.pptx');layout=r.slides[1].slide_layout;original=list(r.slides._sldIdLst);refs=json.loads((P/'references.json').read_text());summary=json.loads((R/'summary.json').read_text());cfg=json.loads((P/'protocol.json').read_text());new=[]
def text(s,x,y,w,h,content,size=22,bold=False,color='000000',after=8):
 sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame;tf.word_wrap=True;tf.margin_left=0;tf.margin_right=0;tf.margin_top=0;tf.margin_bottom=0
 for i,line in enumerate(content.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.space_after=Pt(after);p.line_spacing=1.06;rr=p.add_run();rr.text=line;rr.font.name='Calibri';rr.font.size=Pt(size);rr.font.bold=bold;rr.font.color.rgb=RGBColor.from_string(color)
  for tag in ['ea','cs']:
   e=OxmlElement('a:'+tag);e.set('typeface','Calibri');rr._r.get_or_add_rPr().append(e)
 return sh
def slide(title,source):
 s=r.slides.add_slide(layout)
 for sh in list(s.shapes):
  if sh!=s.shapes.title:sh.element.getparent().remove(sh.element)
 s.shapes.title.text=title;text(s,12.35,7.13,.6,.3,'0',14,color='FFFFFF');text(s,.73,6.83,11.9,.23,source,10,color='555555',after=0)
 s.notes_slide.notes_text_frame.text='v3 / 2026-10-28. '+source;new.append(r.slides._sldIdLst[-1]);return s
def body(s,lines,y=1.14,size=23,h=5.4):return text(s,.8,y,11.75,h,'\n'.join(lines),size,after=13)
def image(s,p,x=.8,y=3.75,w=11.75,h=2.65):
 p=Path(p);iw,ih=Image.open(p).size;sc=min(w/iw,h/ih);pw=iw*sc;ph=ih*sc
 return s.shapes.add_picture(str(p),Inches(x+(w-pw)/2),Inches(y+(h-ph)/2),width=Inches(pw),height=Inches(ph))
def eq(s,n,y=5.85,h=.65):
 sh=image(s,A/(n+'_eq.png'),y=y,h=h);sh.name='LaTeX: '+n;src=(A/(n+'.tex')).read_text();sh._element.nvPicPr.cNvPr.set('descr',src);s.notes_slide.notes_text_frame.text+='\nLaTeX:\n'+src

def table(s,heads,rows,y=2.1,h=3.3,size=19,widths=None):
 t=s.shapes.add_table(len(rows)+1,len(heads),Inches(.8),Inches(y),Inches(11.75),Inches(h)).table
 if widths:
  for col,w in zip(t.columns,widths):col.width=Inches(w)
 for i,row in enumerate([heads]+rows):
  for j,v in enumerate(row):
   c=t.cell(i,j);c.text=str(v);c.margin_top=Inches(.045);c.margin_bottom=Inches(.02);c.fill.solid();c.fill.fore_color.rgb=RGBColor.from_string('254B95' if i==0 else ('F0F3F8' if i%2 else 'FFFFFF'))
   for p in c.text_frame.paragraphs:p.font.name='Calibri';p.font.size=Pt(size);p.font.bold=i==0;p.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else '000000')
 return t
# Correct scope wording on original slides without changing their visual template.
replaces={
'이번 파일럿에서는 지면 평면과 공통 UV를 고정한 상태에서 관측 선택 S를 바꾸며, 차량 제거와 배경 보존 사이의 차이를 확인함.':'같은 장소의 연도별 사진을 평면 좌표에 맞춘 뒤, 차량 관측을 제외하는 S 규칙을 비교함. 차량 흔적과 차선·지면 디테일을 함께 평가함.',
'이번 파일럿에서는 지면 평면과 공통 UV를 고정한 상태에서 관측 선택 S를 바꾸며, 차량 제거와 배경 보존 사이의 차이를 확인함.':'같은 장소의 연도별 사진을 평면 좌표에 맞춘 뒤, 차량 관측을 제외하는 S 규칙을 비교함. 차량 흔적과 차선·지면 디테일을 함께 평가함.',
'메쉬와 투영 관계를 고정하고':'평면 좌표와 카메라 투영 관계를 고정하고',
'지면을 평면 메쉬로 두고':'지면을 z=0 평면으로 가정하고',
'동적 픽셀에서는 S → 0, 정적 픽셀에서는 S → 1이 되도록 설정함.':'이진 선택은 동적 픽셀 S=0, 정적 픽셀 S=1을 목표로 함. 학습형 soft 가중치는 1보다 클 수도 있음.',
'이번 범위: 평면 메쉬의 텍스처 베이킹 파일럿':'이번 범위: 평면 투영식으로 모사한 텍스처 베이킹 파일럿'
}
for s in r.slides:
 for sh in s.shapes:
  if sh.has_text_frame:
   for p in sh.text_frame.paragraphs:
    for rr in p.runs:
     for a,b in replaces.items():rr.text=rr.text.replace(a,b)
# Exact original Goal wording has longer middle phrase; make the relevant paragraph explicit.
for sh in r.slides[0].shapes:
 if sh.has_text_frame:
  for p in sh.text_frame.paragraphs:
   if '이번 파일럿에서는' in p.text:
    for k,rr in enumerate(p.runs):rr.text='이번 파일럿은 지면을 평면으로 가정한 베이킹 모사 실험. 같은 위치의 연도별 사진에서 어떤 관측을 채택할지 바꿔 차량 흔적과 배경 보존을 비교함.' if k==0 else ''
# Planar scope slide, inserted after original Goal.
s=slide('Pilot : What is the geometry?','구현: baking_pilot/pilot.py:65–83, render():107 | 실제 mesh 재구성·삼각형 rasterization 없이 homography로 모사')
body(s,['사진으로 mesh를 복원한 실험이 아님. 120m×120m 지면을 높이 z=0인 평면으로 가정함.','정사영상의 가로·세로 좌표를 평면 x·y에 선형 대응시켜, 같은 지면 위치가 같은 Texture 위치를 사용하도록 함.','평면을 바라보는 가상 카메라 4개(0° / 20° / 35° / 45°)의 homography를 계산하고, grid sampling으로 예측 관측을 생성함.','정점·삼각형 mesh 파일 생성, 실제 3D geometry 추정과 parallax 검증은 수행하지 않음. 비교 대상은 이 조건에서의 관측 선택 규칙 S임.'],size=23,h=4.2)
eq(s,'planar',y=5.75,h=.75);scope=new[-1]
refslides=[]
ppt_summary={
'on_the_go':['DINOv2 특징→MLP로 픽셀별 불확실성 β를 예측하고, 유사 특징 사이의 일관성을 규제함.','변형 SSIM의 밝기·대비·구조 오차 곱으로 β를 학습해, 색이 비슷한 distractor를 구분함.','1/(2β²) 가중 RGB 손실로 NeRF를 학습. Gradient를 분리하고 dilated patch sampling을 사용함.'],
'robustnerf':['정적 NeRF의 큰 photometric residual을 outlier로 취급하며, 동적 물체를 별도 3D 표현으로 만들지 않음.','잔차 분위수와 주변 픽셀·패치의 공간 일관성으로 inlier 마스크를 반복 갱신함.','채택된 관측으로 정적 장면을 학습. 차량·사람 클래스나 semantic detector를 사전에 지정하지 않음.'],
'spotless':['Stable Diffusion 특징에서 의미가 비슷한 픽셀을 묶어 distractor의 공간 범위를 찾음.','Clustering 투표(SLS-agg) 또는 약한 잔차 label로 학습한 inlier MLP(SLS-mlp)를 사용함.','점진적 masking, 잔차 분포 추적, utilization pruning으로 robust mask와 3DGS 학습의 충돌을 줄임.'],
'wildgaussians':['입력 DINO 특징에서 affine+softplus로 σ를 예측. 입력/렌더 DINO cosine 차이로 학습함.','큰 gradient가 densification을 방해하지 않도록 1/(2σ²)>1인 위치만 binary mask로 채택함.','이미지별·Gaussian별 appearance embedding과 색변환을 학습해 조명 변화와 transient를 함께 처리함.'],
'gs_w':['고유 intrinsic appearance와 조명·날씨에 따른 dynamic appearance를 분리하고 특징을 적응 샘플링함.','Transient object에는 U-Net의 2D visibility map VM을 사용해 관측 손실을 가중함.','VM을 1로 유도하는 정규화로 전부 가리는 퇴화를 막음. Dynamic appearance는 이동 물체와 다른 개념임.'],
'robust_gs':['주제는 동적 물체 제거가 아니라 포즈 오차·motion blur·defocus blur·색 불일치에 강한 3DGS임.','카메라 포즈의 Gaussian 분포로 motion blur와 포즈를 보정하고, 투영 covariance로 defocus를 다룸.','이미지별 affine 색변환을 사용함. 차량을 식별해 제외하는 S나 transient 분리 모듈은 없음.'],
'robust_3dgs':['RGB 채널별 centered residual과 이웃·패치 문맥으로 distractor 후보 단서를 만듦.','고정 임계값 대신 logistic regression의 결정 경계와 마스킹 면적 정규화로 채택 여부를 학습함.','SAM object mask로 조각난 픽셀 판정을 object-aware mask로 정리해 정적 3DGS 학습에 적용함.']}
ppt_relation={
'on_the_go':'S를 연속 uncertainty 가중치로 학습하는 근거. 이번에는 고정 warm-up에서 이 부분만 차용했고 원 논문의 공동 최적화는 실행하지 않음.',
'robustnerf':'기존 Robust형 S의 원형. 반복 차량의 잔차가 작아지는 실패는 낮은 잔차가 항상 정적 배경을 의미하지 않음을 보여줌.',
'spotless':'특징 기반 영역 선택·약한 supervision은 이식 가능. 이번 DINO 실험은 SD 특징과 pruning을 쓰는 SpotLessSplats 재현이 아님.',
'wildgaussians':'조명 변화와 차량 변화를 분리할 근거. Gaussian densification이 없는 고정 Texture에서 동일 σ의 soft/hard 선택만 비교함.',
'gs_w':'학습된 2D visibility를 S로 쓸 수 있으나 전부 제외하는 퇴화를 막아야 함. 파일럿의 전역 affine은 GS-W의 국소 appearance보다 제한적임.',
'robust_gs':'정합·블러·색 차이도 잔차를 키운다는 보조 레퍼런스. 동적 제거를 의도한 경우 다음 Robust 3DGS와 구분해야 함.',
'robust_3dgs':'동적 제거 문맥에서는 이 논문이 직접 관련됨. 이번에는 SAM 영역 투표만 차용했고 logistic boundary와 원래 IoU 절차는 생략함.'}
titles={'on_the_go':'Ref : NeRF On-the-go','robustnerf':'Ref : RobustNeRF','spotless':'Ref : SpotLessSplats','wildgaussians':'Ref : WildGaussians','gs_w':'Ref : GS-W','robust_gs':'Ref : Robust Gaussian Splatting','robust_3dgs':'Ref : Robust 3DGS (distractors)'}
for a in refs:
 s=slide(titles[a['key']],a['name']+' | '+a['section']+' | '+a['paper'])
 text(s,.8,1.15,6.05,3.35,'\n'.join(str(i+1)+'. '+q for i,q in enumerate(ppt_summary[a['key']])),20,after=10)
 image(s,A/(a['key']+'_paper.png'),x=7.1,y=1.4,w=5.4,h=2.9)
 text(s,7.1,4.32,5.4,.25,'원 논문 피겨 인용 · '+a['year'],11,color='555555')
 text(s,.8,4.84,11.75,.85,'우리 실험 : '+ppt_relation[a['key']],19,after=0)
 eq(s,'ref_'+a['key'],y=5.91,h=.62)
 s.notes_slide.notes_text_frame.text+='\n3줄 요약:\n'+'\n'.join(a['summary'])+'\n파일럿 관련성: '+a['relation']+'\n실행 범위: '+a['experiment']+'\n공식 구현/프로젝트: '+str(a['code'])
 refslides.append(new[-1])
s=slide('References : What transfers?','원문 Method를 S의 관점에서 비교 | 논문 전체 구현과 파일럿으로 차용한 부분을 구분')
table(s,['연구','S / 핵심 단서','고정 Texture에 이식','v3 상태'],[
 ['On-the-go','DINO + 변형 SSIM','연속 uncertainty 가중','고정 S 차용'],['RobustNeRF','RGB 잔차 + 문맥','반복 binary masking','기존 결과'],['SpotLessSplats','SD 특징 + 약한 label','특징 영역 / inlier MLP','미실행'],['WildGaussians','DINO cosine + σ','soft / hard S 비교','고정 S 차용'],['GS-W','U-Net visibility','visibility + 정규화','미실행'],['Robust GS','블러·포즈·색 보정','오차 원인 분리','dyn S 없음'],['Robust 3DGS','학습 경계 + SAM','잔차의 object-aware 보완','SAM 투표만']],y=1.3,h=4.15,size=18,widths=[2.5,3.1,3.6,2.55])
text(s,.8,5.75,11.75,.75,'Gaussian 생성·분할·pruning과 appearance 모델은 원 논문 성능의 일부임.\nS만 평면 베이킹으로 옮긴 결과를 원 논문의 성능으로 표기하지 않음.',21)
refslides.append(new[-1])
# New experiment group, inserted after original full-scene result.
exp=[]
s=slide('v3 : Frozen-S transfer setup','실험 전 고정: reference_study_v3/protocol.json | 4 methods × 2 scenes × 3 seeds = 24 baking runs')
body(s,['공식 DINOv2 ViT-S/14 특징과 저장된 SAM 분리 영역을 사용. 학습 영상은 2021–2025의 20개 합성 관측이며 2026은 평가에만 사용함.','공통 L2 warm-up(600회)의 예측으로 uncertainty 또는 영역 선택을 만든 뒤 S를 고정. 원 논문의 온라인 공동 갱신과 다른 이식 실험임.','규칙 : DINO-SSIM soft / DINO-cosine soft / 동일 σ의 cosine hard / SAM-region median. Stable Diffusion SLS와 GS-W U-Net은 미실행.','새 baking은 평균 초기값, 512²→1024², Adam 600회, batch 4, seed42/43/44. 모든 사전 지정 방법과 실패를 보고함.','Soft S는 1보다 클 수 있음. 분모가 valid 픽셀 수이므로 S의 평균 크기가 데이터항과 색 정규화의 상대 크기도 바꿀 수 있음.'],size=22,h=4.5)
eq(s,'new_uncertainty',y=5.88,h=.65);exp.append(new[-1])
def g(scene,m):return next(z for z in summary if z['scene']==scene and z['method']==m)
labels={'l2':'L2 (기존)','robust':'Robust형 (기존)','uv_shadow_guard':'UV + shadow (기존)','dino_ssim_soft':'DINO-SSIM soft','dino_cosine_soft':'DINO-cosine soft','dino_cosine_hard':'DINO-cosine hard','sam_region_median':'SAM-region median'}
s=slide('v3 : Quantitative results','새 방법은 원 논문 전체 재현이 아님 | PSNR: 3 seeds 평균 ± 표본표준편차 | 기존 direct fusion은 1회')
body(s,['이번 고정 S 이식에서는 기존 UV + shadow guard를 넘지 못함. Masking 규칙을 옮기는 것만으로 원 논문의 이득이 유지되지는 않았음.'],size=21,h=.8)
rows=[]
for m in labels:
 b=g('industrial_b',m);c=g('industrial_c',m)
 def val(z):return f"{z['PSNR_occupied3_4']:.2f}"+(f" ± {z['PSNR_occupied3_4_std']:.2f}" if z['PSNR_occupied3_4_std'] is not None else '')
 rows.append([labels[m],val(b),val(c),f"{b['LPIPS_object_crop_mean']:.3f} / {c['LPIPS_object_crop_mean']:.3f}"])
table(s,['규칙','B 고점유 PSNR ↑','C 고점유 PSNR ↑','B / C LPIPS ↓'],rows,y=2.14,h=3.65,size=18,widths=[3.6,2.7,2.7,2.75])
text(s,.8,5.93,11.75,.68,'Hard mask의 일부 지표 개선은 아래 관측 지지 진단과 함께 해석해야 함.\nGaussian 모델 또는 전체 NeRF 사이의 순위를 의미하는 표가 아님.',19);exp.append(new[-1])
s=slide('v3 : Mask diagnostics','근거: results/*/mask_diagnostics.json, summary.json | 채택률은 카메라 valid 영역, 무지지는 지면 ROI에서 계산')
rows=[]
for m in cfg['methods']:
 b=json.loads((R/'industrial_b/mask_diagnostics.json').read_text())['stats'][m];c=json.loads((R/'industrial_c/mask_diagnostics.json').read_text())['stats'][m]
 rows.append([labels[m],f"{100*b['positive_fraction']:.1f}% / {100*c['positive_fraction']:.1f}%",f"{b['mean_weight']:.3f} / {c['mean_weight']:.3f}",f"{100*g('industrial_b',m)['unsupported_roi']:.1f}% / {100*g('industrial_c',m)['unsupported_roi']:.1f}%"])
table(s,['규칙','B / C 채택률','B / C 평균 S','B / C 지면 무지지'],rows,y=1.32,h=2.55,size=18,widths=[3.4,2.75,2.75,2.85])
body(s,['SSIM-soft : σ가 하한 0.1에 붙어 S≈50으로 일정해짐. 차량 선택을 학습하지 못한 실패이며 작은 점수 변화는 배율의 영향과도 분리되지 않음.','Cosine-hard : 지면 약 91–95%가 관측 지지를 잃음. 넓은 영역이 평균 초기값에 남을 수 있어 PSNR만으로 제거 성공을 판정할 수 없음.','SAM 투표 역시 지면을 과도하게 제외. 영역 경계를 만드는 것과 정적/동적 영역을 올바르게 판정하는 것은 별개의 문제임.'],y=4.15,size=20,h=2.35);exp.append(new[-1])
s=slide('v3 : What does S select?','실제 입력·저장된 새 S | 2025년, 20° 카메라 | © Beeldmateriaal Nederland / PDOK, CC BY 4.0')
text(s,.8,1.11,11.75,.7,'Soft는 log10 S를 [−1,1]의 공통 색 범위로 표시(파랑: 낮음, 주황: 높음).\nHard/SAM은 흰색이 채택 영역. 회색은 무효 투영이며 학습 가중치를 clip한 것은 아님.',19)
text(s,.8,1.96,11.75,.3,'Industrial B',19,bold=True);image(s,A/'industrial_b_new_masks.png',y=2.27,h=1.94)
text(s,.8,4.43,11.75,.3,'Industrial C',19,bold=True);image(s,A/'industrial_c_new_masks.png',y=4.74,h=1.94);exp.append(new[-1])
s=slide('v3 : Same audit locations','B/C 기존 평가 주석의 첫 OBB, seed42 | crop 선택은 출력 점수와 무관 | © Beeldmateriaal Nederland / PDOK, CC BY 4.0')
body(s,['기존과 새 결과를 같은 위치에서 비교. Hard / 영역 mask의 출력이 덜 오염돼 보여도, 실제로 배경 관측을 사용한 것인지 관측 지지와 함께 확인해야 함.'],size=21,h=.8)
image(s,A/'new_qualitative.png',y=2.0,h=4.13)
text(s,.8,6.36,11.75,.28,'학습 입력 → 실제 2026 배경 → 기존 L2 → cosine soft → cosine hard → SAM 투표 → 기존 UV + shadow',15);exp.append(new[-1])
s=slide('v3 : Interpretation','이번 결과는 이식 설정의 결과 | 마스크 학습은 장면별 seed42 1회, 표준편차는 baking seed만 반영')
body(s,['현재 판단 : 공통 UV가 알려진 이 장면에서는 검출 기반 UV 합성이 여전히 강한 기준선. 새 S 규칙의 우월성은 확인하지 못함.','논문과의 차이 : 고정 warm-up, 작은 항공 차량, feature 리사이즈, 온라인 갱신 생략, Gaussian appearance·density control의 부재가 있음.','후속 1 : 동일 손실·같은 가중치 스케일에서 SSIM target과 cosine target을 비교해, 선택 성능과 gradient 크기 효과를 분리함.','후속 2 : uncertainty를 온라인 갱신하되 지면의 관측 지지 비율을 함께 모니터링. 무지지 영역의 출력은 복원과 분리해서 평가함.','후속 3 : 실제 3D mesh·사선 사진에서 visibility와 parallax를 포함해 검증. Stable Diffusion 특징의 영역 선택과 GS-W VM은 별도 구현·평가가 필요함.'],size=22,h=5.35);exp.append(new[-1])
# Final order retains every original slide and inserts additions at the relevant points.
order=[]
for i,si in enumerate(original):
 order.append(si)
 if i==0:order.append(scope)
 if i==10:order+=refslides
 if i==19:order+=exp
for si in list(r.slides._sldIdLst):r.slides._sldIdLst.remove(si)
for si in order:r.slides._sldIdLst.append(si)
for idx,s in enumerate(r.slides,1):
 for sh in s.shapes:
  if sh.has_text_frame and sh.top>Inches(7) and sh.left>Inches(12) and sh.text.strip().isdigit():
   for p in sh.text_frame.paragraphs:
    for k,rr in enumerate(p.runs):rr.text=str(idx) if k==0 else ''
 s.notes_slide.notes_text_frame.text+='\n수정본 v3: mesh 복원 대신 평면 homography 모사로 범위 명확화. 상세 요약/출처/실험: dynamic_object_baking_references_v3.md'
out=BASE/'dynamic_object_baking_v3.pptx';r.save(out);shutil.copy2(out,'/mnt/data/'+out.name)
(P/'slide_manifest.json').write_text(json.dumps([{'slide':i+1,'title':s.shapes.title.text if s.shapes.title else ''} for i,s in enumerate(r.slides)],ensure_ascii=False,indent=2));print('SAVED',out,'SLIDES',len(r.slides))
