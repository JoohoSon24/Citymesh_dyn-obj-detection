from pathlib import Path
import subprocess,shutil
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE,MSO_CONNECTOR
from pptx.oxml.xmlchemy import OxmlElement
from PIL import Image
P=Path(__file__).parent;BASE=P.parent.parent;r=Presentation(BASE/'dynamic_object_baking_v3.pptx');s=r.slides[6];assert 'HuGS' in s.shapes.title.text
oldnotes=s.notes_slide.notes_text_frame.text
for sh in list(s.shapes):
 if sh!=s.shapes.title:sh.element.getparent().remove(sh.element)
BLUE='254B95';MID='6699CC';LIGHT='EAF0F8';GRAY='E4E7EB';DARK='202020'
def text(x,y,w,h,content,size=20,bold=False,color=DARK,align=None):
 sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame;tf.word_wrap=True;tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
 for i,line in enumerate(content.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.space_after=Pt(5);p.line_spacing=1.05
  if align is not None:p.alignment=align
  rr=p.add_run();rr.text=line;rr.font.name='Calibri';rr.font.size=Pt(size);rr.font.bold=bold;rr.font.color.rgb=RGBColor.from_string(color)
  for tag in ['ea','cs']:
   el=OxmlElement('a:'+tag);el.set('typeface','Calibri');rr._r.get_or_add_rPr().append(el)
 return sh
def rect(x,y,w,h,fill,line=None):
 sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h));sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(fill);sh._element.spPr.append(OxmlElement('a:effectLst'))
 if line:sh.line.color.rgb=RGBColor.from_string(line);sh.line.width=Pt(1.4)
 else:sh.line.fill.background()
 return sh
def arrow(x1,y1,x2,y2):
 sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2));sh.line.color.rgb=RGBColor.from_string('596779');sh.line.width=Pt(1.8);e=OxmlElement('a:tailEnd');e.set('type','triangle');e.set('w','med');e.set('len','med');sh.line._get_or_add_ln().append(e)
 return sh
def grid(x,y,kind):
 # Same two SAM regions, 16 pixels each. Candidate map vs region decision.
 a={0,1,2,3,4,5,6,7,8,9,10,11};b={0,5,10,15};sz=.275
 for row in range(4):
  for col in range(8):
   local=row*4+(col%4);on=(local in (a if col<4 else b)) if kind=='candidate' else col<4
   rect(x+col*sz,y+row*sz,sz-.025,sz-.025,MID if on else 'F2F3F5')
 for reg in [0,1]:
  sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x+reg*4*sz-.035),Inches(y-.035),Inches(4*sz-.005),Inches(4*sz-.005));sh.fill.background();sh._element.spPr.append(OxmlElement('a:effectLst'));sh.line.color.rgb=RGBColor.from_string(BLUE if reg==0 else '68727F');sh.line.width=Pt(2)
 return 8*sz
s.shapes.title.text='S : NeRF-HuGS type'
text(.8,1.12,11.75,.82,'직관 : “정적으로 보이는 픽셀”을 먼저 찾고, SAM으로 나눈 영역마다 다수결함.\n한 영역의 절반 이상이 정적 후보이면 그 영역을 Texture 학습에 사용함.',23)
text(.8,2.16,3.1,.4,'① 정적 후보의 근거',21,True,color=BLUE)
rect(.8,2.78,3.08,.88,LIGHT)
text(.95,2.87,2.79,.72,'여러 영상에서 확인된 점\n→ SAM으로 주변 영역 확장',17)
rect(.8,3.91,3.08,.88,LIGHT)
text(.95,4.0,2.8,.72,'예측과 색이 잘 맞음\n→ warm-up 오차가 작음',17)
# Two sources feed the static candidate map.
arrow(3.92,3.22,4.5,3.65);arrow(3.92,4.36,4.5,3.65);arrow(4.5,3.65,4.96,3.65)
text(5.02,2.16,2.65,.4,'② 정적 후보 지도',21,True,color=BLUE)
text(5.02,2.72,2.55,.65,'두 근거를 합치고\n큰 오차 제외·주변 보정',17)
grid(5.08,3.53,'candidate')
text(5.05,4.84,1.15,.55,'영역 A\n12 / 16',18,True,color=BLUE)
text(6.2,4.84,1.1,.55,'영역 B\n4 / 16',18,True)
arrow(7.57,4.03,9.06,4.03)
text(7.62,3.29,1.4,.62,'영역마다\n절반 이상?',18,True)
text(9.23,2.16,3.1,.4,'③ 영역별 채택 / 제외',21,True,color=BLUE)
text(9.23,2.72,3.1,.65,'A는 채택 / B는 제외\n픽셀별 판정을 영역으로 정리',17)
grid(9.27,3.53,'final')
text(9.24,4.84,1.25,.55,'S = 1\n75% 채택',18,True,color=BLUE)
text(10.5,4.84,1.78,.55,'S = 0\n25% 제외',18,True)
# One simple region-vote equation. H denotes the final static-candidate map.
eq=r'S(m)=\mathbf{1}\!\left[\frac{N_{\mathrm{static}}(m)}{N_{\mathrm{all}}(m)}\ge0.5\right]'
tex=r'\documentclass[border=3pt]{standalone}\usepackage{amsmath,amssymb}\begin{document}$\displaystyle '+eq+r'$\end{document}'
(P/'hugs_simple.tex').write_text(tex);q=subprocess.run(['pdflatex','-interaction=nonstopmode','-halt-on-error','-output-directory',str(P),str(P/'hugs_simple.tex')],stdout=subprocess.PIPE)
if q.returncode:raise RuntimeError(q.stdout.decode()[-2000:])
subprocess.run(['pdftocairo','-png','-singlefile','-transp','-r','300',str(P/'hugs_simple.pdf'),str(P/'hugs_simple')],check=True)
im=Image.open(P/'hugs_simple.png');iw,ih=im.size;h=.53;w=h*iw/ih
sh=s.shapes.add_picture(str(P/'hugs_simple.png'),Inches((13.333-w)/2),Inches(5.58),height=Inches(h));sh.name='LaTeX: simplified HuGS region vote';sh._element.nvPicPr.cNvPr.set('descr',tex)
text(.8,6.15,11.75,.25,'Nstatic : 영역 안 정적 후보 수 / Nall : 영역 전체 픽셀 수. 파랑은 후보(가운데) 또는 채택(오른쪽)을 뜻함.',14,color='555555')
text(.8,6.48,11.75,.3,'도식의 75% / 25%는 이해용 예시. 실제 구현은 마지막에 경계를 조금 축소(erosion)하며, 반복 차량도 정적으로 오인할 수 있음.',13,color='555555')
text(.73,6.83,11.9,.23,'참조: NeRF-HuGS 공식 generate_static_mask.py · pilot.py:hugs() | 파일럿은 3D SfM 대신 2D SIFT 매칭 사용',10,color='555555')
text(12.35,7.13,.6,.3,'7',14,color='FFFFFF')
s.notes_slide.notes_text_frame.text=oldnotes+'\n\nv4: 이해를 위해 두 근거 → 정적 후보 → SAM 영역 투표로 단순화. 32칸/75%/25%는 도식 예시이며 실험 데이터 아님. 단순 수식은 erosion 이전의 영역 채택 조건. 실제 구현: F=SIFT 지지점 SAM 확장, C={eps<=mu}, Q={eps<=q95} union C; H0=(F union C) intersect Q; H=H0 union {B7(H0)>=0.5}; 영역별 H 비율>=0.5를 채택한 뒤 5x5 erosion 및 valid 처리.\n'+tex
out=BASE/'dynamic_object_baking_v4.pptx';r.save(out);shutil.copy2(out,'/mnt/data/'+out.name)
# Standalone edited slide for quick review, retaining its original master.
one=Presentation(out)
for si in list(one.slides._sldIdLst):
 if si!=one.slides._sldIdLst[6]:pass
keep=one.slides._sldIdLst[6]
for si in list(one.slides._sldIdLst):
 if si!=keep:one.part.drop_rel(si.rId);one.slides._sldIdLst.remove(si)
one.save(P/'hugs_page.pptx')
print('Saved v4,',len(r.slides),'slides; changed slide 7 only')
