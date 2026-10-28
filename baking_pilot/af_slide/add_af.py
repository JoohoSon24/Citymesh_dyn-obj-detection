from pathlib import Path
import shutil
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.oxml.xmlchemy import OxmlElement
P=Path(__file__).parent;BASE=P.parent.parent;r=Presentation(BASE/'dynamic_object_baking_v5.pptx');layout=r.slides[1].slide_layout;s=r.slides.add_slide(layout)
for sh in list(s.shapes):
 if sh!=s.shapes.title:sh.element.getparent().remove(sh.element)
s.shapes.title.text='AF'
def text(x,y,w,h,content,size=23,bold=False,color='000000'):
 sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame;tf.word_wrap=True;tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
 for j,line in enumerate(content.split('\n')):
  p=tf.paragraphs[0] if j==0 else tf.add_paragraph();p.line_spacing=1.07;p.space_after=Pt(7);rr=p.add_run();rr.text=line;rr.font.name='Calibri';rr.font.size=Pt(size);rr.font.bold=bold;rr.font.color.rgb=RGBColor.from_string(color)
  for tag in ['ea','cs']:
   el=OxmlElement('a:'+tag);el.set('typeface','Calibri');rr._r.get_or_add_rPr().append(el)
 return sh
text(.8,1.12,11.75,.38,'관찰 : 공통 UV에서 차량을 직접 검출해 제외하는 방식이 가장 효과적이었음.',23,True,color='254B95')
text(.8,1.75,11.75,.82,'이번 평면 조건에서는 원본 정사영상의 차량 마스크로 배경만 모은 UV fusion이 강했고,\nshadow guard를 더했을 때 차량 위치 PSNR·LPIPS가 가장 좋았음.',22)
text(.8,2.87,11.75,.38,'해석 : 시간적 움직임보다 관측의 불일치 또는 의미 정보를 이용하는 문제에 가까움.',23,True,color='254B95')
text(.8,3.5,11.75,.95,'사진의 촬영 시점·간격이 불규칙해 연속 프레임처럼 시간적 연속성을 기대하기 어려움.\n검토한 Ref도 시간 순서 자체보다 입력과 재구성 예측의 차이·다중 시점 일관성을 보거나,\nsemantic / 특징·영역 정보를 활용해 신뢰할 관측을 선택하는 방향을 사용함.',21)
text(.8,4.85,11.75,.38,'아쉬움 : 선택 규칙에 휴리스틱이 많이 남아 있음.',23,True,color='254B95')
text(.8,5.48,11.75,.85,'잔차 분위수·주변 필터·마스크 팽창·그림자 밝기 등 고정 설정에 의존함.\nRef의 학습형 모델도 있지만, 이번 단순 이식에서는 이득을 얻지 못함.',20)
text(.8,6.45,11.75,.38,'다음 질문 : 고정 규칙 의존을 줄이고, 배경 관측을 남기는 S를 학습할 수 있을까?',19)
text(12.35,7.13,.6,.3,'26',14,color='FFFFFF')
s.notes_slide.notes_text_frame.text='AF: user-requested discussion slide. Observation limited to the current planar homography pilot; not a claim that all references are class-specific semantic removal or simple image-pair differencing. UV masks are detector outputs on original aligned ortho imagery, then fused in common UV. Best vehicle PSNR/LPIPS among shown methods: UV + shadow guard. References include learned uncertainty/visibility and semantic feature/region evidence; input-versus-reconstruction residual is distinguished from direct image-to-image subtraction. Training captures are separated by years, not consecutive video frames. Heuristic concerns refer to current implementation choices. Sources: PROJECT.md, reference_study_v3/REPORT.md, protocol.json, results/summary.json. No new experiments.'
# Preserve the two final Result pages as the final two slides.
si=r.slides._sldIdLst[-1];r.slides._sldIdLst.remove(si);r.slides._sldIdLst.insert(25,si)
for i,sl in enumerate(r.slides,1):
 for sh in sl.shapes:
  if sh.has_text_frame and sh.top>Inches(7) and sh.left>Inches(12) and sh.text.strip().isdigit():
   for pa in sh.text_frame.paragraphs:
    for k,rr in enumerate(pa.runs):rr.text=str(i) if k==0 else ''
out=BASE/'dynamic_object_baking_v6.pptx';r.save(out);shutil.copy2(out,'/mnt/data/'+out.name)
# Standalone AF preview.
one=Presentation(out);keep=one.slides._sldIdLst[25]
for item in list(one.slides._sldIdLst):
 if item!=keep:one.part.drop_rel(item.rId);one.slides._sldIdLst.remove(item)
one.save(P/'AF.pptx')
print('Saved 28 slides; AF is page 26; Result pages remain 27–28.')
