from pathlib import Path
from copy import deepcopy
import json,shutil
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE,MSO_CONNECTOR
from pptx.oxml.xmlchemy import OxmlElement
from pptx.opc.packuri import PackURI
from PIL import Image
P=Path(__file__).parent;BASE=P.parent.parent;A=P/'assets';ST=P.parent/'reference_study_v3';OLD=Path('/mnt/data/baking_revision/assets');r=Presentation(BASE/'dynamic_object_baking_v4.pptx');layout=r.slides[1].slide_layout
ids=list(r.slides._sldIdLst);keepidx=[0,2,20,1,21,3,22];keep=[ids[i] for i in keepidx]
for i,si in enumerate(ids):
 r.slides._sldIdLst.remove(si)
 if i not in keepidx:r.part.drop_rel(si.rId)
for si in keep:r.slides._sldIdLst.append(si)
for i,s in enumerate(r.slides,1):s.part.partname=PackURI(f"/ppt/slides/slide{i}.xml")
# Remove on-slide source lines; source details remain in speaker notes.
for s in r.slides:
 for sh in list(s.shapes):
  if sh.has_text_frame and ('참조:' in sh.text or '참조 :' in sh.text or (sh.top>=Inches(6.75) and sh.top<Inches(7.0))):
   s.notes_slide.notes_text_frame.text+='\n'+sh.text;sh.element.getparent().remove(sh.element)
 # Earlier rendering text doesn't need per-version narration.
 for sh in s.shapes:
  if sh.has_text_frame:
   for p in sh.text_frame.paragraphs:
    for rr in p.runs:rr.text=rr.text.replace('v3','').replace('v4','')

def text(s,x,y,w,h,content,size=22,bold=False,color='000000',after=7):
 sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame;tf.word_wrap=True;tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
 for i,line in enumerate(content.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.space_after=Pt(after);p.line_spacing=1.05;rr=p.add_run();rr.text=line;rr.font.name='Calibri';rr.font.size=Pt(size);rr.font.bold=bold;rr.font.color.rgb=RGBColor.from_string(color)
  for tag in ['ea','cs']:
   e=OxmlElement('a:'+tag);e.set('typeface','Calibri');rr._r.get_or_add_rPr().append(e)
 return sh

def slide(title,notes=''):
 s=r.slides.add_slide(layout)
 for sh in list(s.shapes):
  if sh!=s.shapes.title:sh.element.getparent().remove(sh.element)
 s.shapes.title.text=title;text(s,12.35,7.13,.6,.3,'0',14,color='FFFFFF');s.notes_slide.notes_text_frame.text=notes;return s

def img(s,path,x=.8,y=4.2,w=11.75,h=2.2):
 path=Path(path);iw,ih=Image.open(path).size;sc=min(w/iw,h/ih);pw=iw*sc;ph=ih*sc
 sh=s.shapes.add_picture(str(path),Inches(x+(w-pw)/2),Inches(y+(h-ph)/2),width=Inches(pw),height=Inches(ph));return sh

def equation(s,path,y=2.74,h=.64):
 path=Path(path);sh=img(s,path,y=y,h=h);sh.name='LaTeX: '+path.stem
 tex=path.with_name(path.stem.replace('_eq','')+'.tex')
 if tex.exists():s.notes_slide.notes_text_frame.text+='\nLaTeX:\n'+tex.read_text()
 return sh

def bodyreplace(s,lines,size=23):
 for sh in list(s.shapes):
  if sh.has_text_frame and Inches(.95)<sh.top<Inches(6.8):sh.element.getparent().remove(sh.element)
 text(s,.8,1.13,11.75,5.6,'\n'.join(lines),size,after=13)
bodyreplace(r.slides[2],[
 '데이터 : PDOK 네덜란드 항공 정사영상. 실제 차량이 있는 산업 주차장 B / C 두 장면을 사용함.',
 '연도 : 2021–2025의 5개 연도를 학습에 사용하고, 2026은 차량 위치의 실제 빈 배경을 확인하는 평가 영상으로 유지함.',
 '장면 크기 : 각각 120m×120m. 모든 방법에 동일한 정합·색 정규화 영상과 고정 지면 좌표를 사용함.',
 '고점유 조건 : 검출 경험 픽셀 중 3년 이상 차량이 점유한 비율은 B 68.1%, C 63.0%. 배경이 소수 연도에만 보이는 상황을 포함함.',
 '비교 원칙 : 학습된 결과 Texture를 같은 차량 위치에서 확인하고, 차량 흔적 감소와 정적 지면 보존을 함께 평가함.'])
bodyreplace(r.slides[6],[
 '차량 위치 PSNR ↑ : 학습 5회 중 3–4회 점유된 고정 OBB 21개(B4 / C17)를 2026년 실제 배경과 비교함.',
 'LPIPS ↓ : 전체 OBB 33개(B16 / C17)의 고정 80×80 crop에서 평가. 폴리곤 바깥은 참조 영상으로 채워 비교함.',
 '정적 PSNR ↑ : 학습·평가 영상 모두 검출 차량이 없고 연도별 색변화가 적은 지면 대조영역에서 계산함.',
 '정성 비교 : B/C의 평가 주석 첫 OBB를 모든 방법에 똑같이 사용함. 결과 그림은 seed42이며, 출력이 좋은 위치를 골라 바꾸지 않음.',
 '학습형 지표는 3 seeds 평균 ± 표본표준편차. 직접 중앙값·UV 합성은 결정적 1회. 연도별 배경 변화와 정합오차가 남는 proxy 평가임.'])

config=json.loads((P/'methods.json').read_text());methods=config['methods'];labels=config['labels'];summary=json.loads((ST/'results/summary.json').read_text());meta=json.loads((A/'manifest.json').read_text())
DATA={
'l2':('Baseline : L2',[
 '모든 유효 관측을 같은 가중치 S=1로 사용. 단순한 한 texel에서는 관측의 산술평균을 맞추는 방향임.',
 '차량이 여러 연도에 반복되면 차량 색도 Texture에 반영되어 잔상·중첩 흔적이 남을 수 있음.',
 '아래 가운데는 실제 학습한 최종 Texture. 동일 위치의 차량 입력과 2026 빈 배경을 함께 비교함.'],OLD/'baseline_eq.png'),
'median':('Baseline : Temporal median',[
 '같은 지면 위치의 5개 학습 연도 색을 채널별 중앙값으로 합침. 별도의 역렌더링 최적화는 하지 않음.',
 '배경이 다수로 관측되면 차량을 줄일 수 있지만, 차량이 3–4회 반복되면 차량 색을 선택할 수 있음.',
 '관측 선택을 학습하는 방법이 아니라, 다중 시기 합성의 직접 기준선임.'],OLD/'median_eq.png'),
'block':('S : Block-NeRF type',[
 'DOTA YOLO11s-OBB로 입력별 차량 영역을 검출하고, 5×5 팽창한 제외 마스크 D를 생성함.',
 '차량으로 검출한 픽셀을 S=0으로 두고 나머지를 baking. 검출 누락과 차량 OBB 밖의 그림자는 남을 수 있음.',
 '원 논문의 semantic 제외 아이디어를 차용한 평면 baking이며, 아래는 마스크가 아닌 최종 학습 Texture임.'],OLD/'block_eq.png'),
'hugs':('S : NeRF-HuGS type',[],P.parent/'hugs_explanation_v4/hugs_simple.png'),
'robust':('S : RobustNeRF type',[
 '현재 예측과의 잔차가 작은 픽셀을 inlier로 보고, 주변 픽셀과 패치의 일관성으로 S를 갱신함.',
 '차량이 여러 번 반복돼 예측에 흡수되면 잔차가 작아져, 정적 배경처럼 계속 채택될 수 있음.',
 '원 논문 전체 NeRF 대신 동일한 평면 Texture에 robust 선택 규칙을 적용한 결과임.'],OLD/'robust_eq.png'),
'geo_mask':('S : Geometry mask',[
 '원본 정사영상에서 검출한 차량 마스크를 공통 UV에 두고, 알려진 homography로 각 카메라에 투영함.',
 '같은 연도의 네 시점이 같은 차량 검출을 공유해 시점별 독립 검출의 차이를 줄임.',
 '공통 평균 초기값과 600회 baking은 유지. 아래 가운데에서 실제 차량 잔상과 지면 질감을 확인함.'],OLD/'geo_eq.png'),
'uv_fusion':('Direct UV fusion',[
 '공통 UV에서 차량으로 검출되지 않은 연도만 선택해, texel별 RGB 중앙값으로 직접 합성함.',
 '배경 관측이 소수여도 활용할 수 있음. 모든 연도에 가려진 위치는 최근접 초기화와 조화 보간으로 채움.',
 '최종 출력은 공통 512² 표현을 거쳐 1024²로 만듦. 미관측 위치의 보간값은 실제 배경 관측과 구분해야 함.'],OLD/'fusion_eq.png'),
'uv_shadow_guard':('UV fusion + Shadow guard',[
 'UV fusion에서 차량 주변 반경 18 texel의 어두운 관측도 추가 제외. 비차량 주변 평균 밝기의 0.72 미만을 그림자 후보로 봄.',
 '같은 배경 합성과 미관측부 보간을 적용. 원래 어두운 포장·표식도 제외할 수 있는 밝기 휴리스틱임.',
 '아래 가운데는 shadow guard까지 적용한 최종 Texture. 입력이나 빨간 마스크가 아니라 제거 후 결과를 비교함.'],OLD/'shadow_eq.png'),
'geo_fusion':('Geometry + Fusion init',[
 'Geometry mask는 유지하고, 평균 대신 UV fusion 결과를 Texture 초기값으로 사용함.',
 '관측이 없는 위치에 배경 추정값을 제공한 뒤 동일한 600회 baking을 수행함.',
 '초기값의 이득이 최적화 뒤에도 남는지 비교하는 변형. 결과는 직접 UV fusion과 다를 수 있음.'],OLD/'geo_eq.png'),
'geo_smooth':('Geometry + Spatial prior',[
 'Geometry mask와 fusion 초기값에, 관측이 적은 위치일수록 강해지는 TV와 미관측부 anchor를 추가함.',
 '불안정한 영역을 매끄럽게 하고 초기 추정값을 유지하려는 목적. 작은 차선·Texture 디테일도 흐려질 수 있음.',
 '같은 데이터와 600회 baking을 사용하므로 아래 최종 결과에서 prior의 효과를 확인함.'],OLD/'prior_eq.png'),
'dino_ssim_soft':('DINO : SSIM uncertainty',[
 'On-the-go의 특징→불확실성→가중치 아이디어를 차용. DINO 특징으로 변형 SSIM 오차를 설명하는 σ를 학습함.',
 '공통 L2 warm-up에서 S를 만든 뒤 고정하고 baking. 온라인 공동 학습·dilated sampling을 포함한 원 논문 전체 재현은 아님.',
 '이 설정에서는 σ가 하한에 붙어 S≈50으로 일정해짐. 선택 능력을 얻지 못했고 L2와의 차이는 손실 배율 영향도 포함함.'],ST/'assets/ref_on_the_go_eq.png'),
'dino_cosine_soft':('DINO : Cosine soft S',[
 '입력과 L2 warm-up 렌더의 DINO cosine 차이로 불확실성 σ를 학습. 외관 차이보다 특징 차이에 반응시키려는 설정임.',
 'S=1/(2σ²)의 연속 가중치로 baking. 다음 hard 방법과 동일한 σ를 사용해 최종 선택 방식만 비교함.',
 'S가 양수여도 올바른 배경 관측을 뜻하지는 않음. 큰 soft 가중치는 데이터항과 색 정규화의 상대 배율도 바꿈.'],ST/'assets/new_uncertainty_eq.png'),
'dino_cosine_hard':('DINO : Cosine hard S',[
 'WildGaussians의 uncertainty→binary mask 규칙을 차용. 앞 페이지와 같은 σ에서 1/(2σ²)>1인 관측만 채택함.',
 'Gaussian appearance·densification·온라인 갱신은 제외한 고정 S 이식 실험임.',
 '지면의 약 91–95%가 관측 지지를 잃었음. 최종 결과가 덜 오염돼 보여도 평균 초기값에 남은 영역과 분리해 해석해야 함.'],ST/'assets/ref_wildgaussians_eq.png'),
'sam_region_median':('S : SAM-region vote',[
 'Warm-up 잔차가 중앙값보다 큰 후보를 만들고, SAM 분리 영역에서 제외 후보 비율이 40%를 넘으면 영역 전체를 제외함.',
 'Robust 3D Gaussian Splatting의 object-aware 취지를 단순 차용. 학습형 logistic boundary와 원래 SAM IoU 절차는 생략함.',
 '영역 경계를 보존해도 정적/동적 구분이 보장되지는 않음. 이번에는 넓은 지면을 제외하는 실패가 나타남.'],ST/'assets/ref_robust_3dgs_eq.png')}

for _m,_file in [('sam_region_median','sam_vote_eq.png'),('uv_shadow_guard','shadow_simple_eq.png')]:
 _t,_l,_e=DATA[_m];DATA[_m]=(_t,_l,A/_file)

def result_row(s,m):
 year=meta['industrial_c']['annotation']['year']
 labs=[f'차량이 있는 입력 ({year})','최종 베이킹 Texture','실제 빈 배경 (2026)'];files=['industrial_c_input.png','industrial_c_'+m+'.png','industrial_c_reference.png']
 for j,(lab,fn) in enumerate(zip(labs,files)):
  x=.8+j*4.0;text(s,x,3.92,3.8,.35,lab,18,bold=True,color='254B95' if j==1 else '000000')
  pic=img(s,A/fn,x=x,y=4.38,w=3.75,h=2.15);pic.name='Actual final baking output: '+m if j==1 else lab
 text(s,.8,6.72,11.75,.25,'동일한 C 장면 평가 위치 · seed42 · 중앙 패널은 각 방법의 최종 산출 Texture',12,color='555555')
 s.notes_slide.notes_text_frame.text+='\nFinal output path: '+meta['industrial_c']['textures'][m]+'\nCrop: '+str(meta['industrial_c']['crop'])+'\nInput year: '+str(year)+'; reference 2026; same first fixed OBB for every method. Imagery © Beeldmateriaal Nederland / PDOK, CC BY 4.0.'

def rect(s,x,y,w,h,fill):
 sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h));sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(fill);sh.line.fill.background();sh._element.spPr.append(OxmlElement('a:effectLst'));return sh

def arrow(s,x,y,w):text(s,x,y,w,.4,'→',25,bold=True,color='596779')
def hugs(s):
 text(s,.8,1.12,11.75,.7,'직관 : 여러 영상에서 반복 확인되거나 예측 오차가 작은 픽셀을 정적 후보로 봄.\nSAM으로 나눈 한 영역의 절반 이상이 정적 후보이면 그 영역을 Texture 학습에 사용함.',21)
 rect(s,.8,2.06,2.9,.43,'EAF0F8');text(s,.93,2.11,2.67,.32,'여러 영상에서 확인된 점',16)
 rect(s,.8,2.65,2.9,.43,'EAF0F8');text(s,.93,2.7,2.67,.32,'예측과 색이 잘 맞는 픽셀',16)
 arrow(s,3.79,2.48,.4);text(s,4.28,2.02,1.8,.3,'정적 후보',17,True)
 for ox,mode in [(4.3,'candidate'),(9.8,'final')]:
  for y in range(4):
   for x in range(8):
    n=y*4+x%4;on=(n<12 if x<4 else n in [0,5,10,15]) if mode=='candidate' else x<4
    rect(s,ox+x*.145,2.51+y*.145,.132,.132,'6699CC' if on else 'E4E7EB')
 arrow(s,5.62,2.48,.5);text(s,6.15,2.03,3.2,.87,'SAM 영역마다 후보 비율 확인\n75% : 영역 전체 채택\n25% : 영역 전체 제외',17)
 arrow(s,9.22,2.48,.4);text(s,9.8,2.02,2.2,.3,'최종 S',17,True)
 text(s,.8,3.27,11.75,.25,'격자는 이해용 예시. 실제 구현은 큰 오차 제외·주변 보정 후 영역 투표를 하고 마지막에 경계를 조금 축소함.',12,color='555555')
 equation(s,P.parent/'hugs_explanation_v4/hugs_simple.png',y=3.45,h=.35)

method_slide_nums={}
for m in methods:
 title,lines,eqp=DATA[m];s=slide(title,'Method: '+m+'; detailed scope and paper sources: dynamic_object_baking_references_v3.md. '+str(eqp));method_slide_nums[m]=len(r.slides)
 if m=='hugs':hugs(s)
 else:
  shown=[lines[0],lines[2]] if m in ['dino_ssim_soft','dino_cosine_hard','sam_region_median'] else lines[:2]
  text(s,.8,1.13,11.75,1.95,'\n'.join(shown),21,after=9);s.notes_slide.notes_text_frame.text+='\nFull method scope: '+' '.join(lines);equation(s,eqp,y=3.1,h=.58)
 result_row(s,m)
# Keep non-executed literature visually distinct, with published final renderings.
refs=json.loads((ST/'references.json').read_text())
for key,title in [('spotless','Reference : SpotLessSplats'),('gs_w','Reference : GS-W'),('robust_gs','Reference : Robust GS')]:
 a=next(x for x in refs if x['key']==key)
 if key=='spotless':a['summary']=['Stable Diffusion 특징에서 의미적으로 비슷한 픽셀을 묶어 distractor의 공간 범위를 찾음.','Clustering 투표 또는 약한 잔차 label로 학습한 inlier MLP를 사용해 관측을 선택함.','점진적 masking·잔차 분포 추적·utilization pruning으로 robust mask와 3DGS 학습의 충돌을 줄임.']
 if key=='robust_gs':a['summary']=['주제는 동적 물체 제거가 아니라 포즈 오차·motion/defocus blur·색 불일치에 강한 3DGS임.','카메라 포즈의 Gaussian 분포로 motion blur와 포즈를 보정하고, 투영 covariance로 defocus를 다룸.','이미지별 affine 색변환을 사용함. 차량을 식별해 제외하는 S나 transient 분리 모듈은 없음.']
 s=slide(title,a['paper']+'\n'+str(a['code'])+'\n'+a['experiment']+'\nPublished figure crop, authors retain copyright; not our pilot result.')
 if key=='gs_w':
  text(s,.8,1.13,5.8,4.7,'\n'.join(str(i+1)+'. '+x for i,x in enumerate(a['summary'])),22,after=13)
  img(s,A/'gs_w_paper_result.png',x=7,y=1.2,w=5.4,h=5.3)
  text(s,.8,6.5,11.75,.3,'원 논문의 최종 렌더 비교 · 파일럿에서는 U-Net visibility / Gaussian appearance 미실행',15,color='555555')
 else:
  text(s,.8,1.12,11.75,2.05,'\n'.join(str(i+1)+'. '+x for i,x in enumerate(a['summary'])),21,after=10)
  img(s,A/(key+'_paper_result.png'),y=3.28,h=2.95)
  text(s,.8,6.5,11.75,.3,'원 논문의 최종 렌더 비교 · 자체 실험의 정량 표에는 포함하지 않음',15,color='555555')
  if key=='robust_gs':text(s,.8,6.84,11.75,.22,'이 논문은 블러·포즈·색 보정 연구. 동적 제거용 Robust 3DGS는 SAM-region 페이지와 별도 구분함.',11,color='555555')
# Discussion before final two results; no chronology and no extra result fragments.
s=slide('Analysis / Next steps','All reported own metrics: reference_study_v3/results/summary.json. Frozen-S transfers, not full paper reproductions.')
text(s,.8,1.14,11.75,5.7,'\n'.join([
 '현재 조건에서는 배경 관측을 직접 모으는 UV fusion이 강한 기준선이며, shadow guard가 차량 위치 PSNR과 LPIPS를 추가 개선함.',
 '반복 차량은 작은 잔차를 가질 수 있음. 정적 배경을 선택하는 근거와 영역 경계를 만드는 방법은 분리해서 평가해야 함.',
 '학습된 S가 전부 같아지거나 지면을 지나치게 제외할 수도 있음. 점수뿐 아니라 최종 Texture와 관측 지지 범위를 함께 확인해야 함.',
 '모든 시기에 가려진 배경은 실제 관측으로 복원할 수 없음. 보간·초기값·공간 prior의 추정 영역을 관측 복원과 구분함.',
 '다음 검증 : 온라인 uncertainty 갱신, 가중치 스케일 통제, 실제 3D mesh·사선 사진에서 visibility와 parallax를 포함한 평가.']),23,after=14)
# One qualitative overview covering every executed method shown in the deck.
s=slide('Result : All methods / Qualitative','All 14 shown methods + vehicle inputs + observed 2026 backgrounds. Matched first audit OBBs, seed42, left B/right C. See assets/manifest.json.')
text(s,.8,1.09,11.75,.43,'모든 패널은 동일한 평가 위치 · 각 패널 왼쪽 B / 오른쪽 C · seed42',19)
img(s,A/'all_methods_qualitative.png',y=1.62,h=5.07)
text(s,.8,6.84,11.75,.22,'Input / 2026 실제 배경을 포함해 모든 발표 방법의 최종 Texture를 동일 crop·색 범위로 비교함.',11,color='555555')
# One quantitative overview; no separate chronological experiments.
s=slide('Result : All methods / Metrics','PSNR occupied3_4 mean ± sample std for 3-seed learned methods; LPIPS/static PSNR mean; direct methods single deterministic run. Sources in reference_study_v3/results/summary.json.')
text(s,.8,1.1,11.75,.47,'차량 위치 PSNR ↑ / LPIPS ↓ / 정적 PSNR ↑ · 각 값은 B / C 순서',19)
t=s.shapes.add_table(len(methods)+1,4,Inches(.8),Inches(1.72),Inches(11.75),Inches(4.56)).table
for col,w in zip(t.columns,[3.45,3.2,2.4,2.7]):col.width=Inches(w)
heads=['방법','차량 위치 PSNR ↑ (B / C)','LPIPS ↓ (B / C)','정적 PSNR ↑ (B / C)']
rows=[]
def g(scene,m):return next(z for z in summary if z['scene']==scene and z['method']==m)
for m in methods:
 b=g('industrial_b',m);c=g('industrial_c',m)
 def val(z):return f"{z['PSNR_occupied3_4']:.2f}"+(f"±{z['PSNR_occupied3_4_std']:.2f}" if z['PSNR_occupied3_4_std'] is not None else '')
 rows.append([labels[m],val(b)+' / '+val(c),f"{b['LPIPS_object_crop_mean']:.3f} / {c['LPIPS_object_crop_mean']:.3f}",f"{b['PSNR_static_control_proxy']:.2f} / {c['PSNR_static_control_proxy']:.2f}"])
for i,row in enumerate([heads]+rows):
 for j,v in enumerate(row):
  cell=t.cell(i,j);cell.text=v;cell.margin_top=Inches(.02);cell.margin_bottom=0;cell.margin_left=Inches(.06);cell.margin_right=Inches(.03);cell.fill.solid();cell.fill.fore_color.rgb=RGBColor.from_string('254B95' if i==0 else ('E9F2FB' if row[0]=='UV + shadow guard' else ('F0F3F8' if i%2 else 'FFFFFF')))
  for p in cell.text_frame.paragraphs:p.font.name='Calibri';p.font.size=Pt(15 if i==0 else 15.5);p.font.bold=i==0 or row[0]=='UV + shadow guard';p.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else '000000')
text(s,.8,6.53,11.75,.37,'学習型は3 seeds平均±標本標準偏差。'.replace('学習型は3 seeds平均±標本標準偏差。','학습형은 3 seeds 평균±표본표준편차. 직접 median·UV 합성은 1회. S 차용 실험을 원 논문 전체 성능으로 해석하지 않음.'),13,color='555555')
# Renumber only; preserve template footer branding and fonts.
for i,s in enumerate(r.slides,1):
 for sh in s.shapes:
  if sh.has_text_frame and sh.left>Inches(12) and sh.top>Inches(7) and sh.text.strip().isdigit():
   for p in sh.text_frame.paragraphs:
    for j,rr in enumerate(p.runs):rr.text=str(i) if j==0 else ''
# Remove any remaining bottom source labels. Citations remain in notes.
for s in r.slides:
 for sh in list(s.shapes):
  if sh.has_text_frame and ('참조:' in sh.text or '참조 :' in sh.text):
   s.notes_slide.notes_text_frame.text+='\n'+sh.text;sh.element.getparent().remove(sh.element)
out=BASE/'dynamic_object_baking_v5.pptx';r.save(out);shutil.copy2(out,'/mnt/data/'+out.name)
(P/'slide_manifest.json').write_text(json.dumps([{'slide':i+1,'title':s.shapes.title.text if s.shapes.title else ''} for i,s in enumerate(r.slides)],ensure_ascii=False,indent=2));(P/'method_slides.json').write_text(json.dumps(method_slide_nums,indent=2))
print('Saved',out,'slides',len(r.slides),'method result rows',len(method_slide_nums))
